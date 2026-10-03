"""富途牛牛與 Firstrade 月結／交易匯出解析。"""

from __future__ import annotations

import csv
import re
from dataclasses import dataclass, field
from datetime import datetime, timedelta
from pathlib import Path
from typing import Any, Iterable


@dataclass
class Trade:
    broker: str
    symbol: str
    side: str  # buy | sell
    quantity: float
    price: float
    trade_date: str  # YYYY-MM-DD
    trade_time: str = ""
    fees: float = 0.0
    amount: float = 0.0
    currency: str = "USD"
    market: str = ""
    source_file: str = ""
    source_row: int = 0


@dataclass
class ParseResult:
    broker: str
    file_name: str
    trades: list[Trade] = field(default_factory=list)
    warnings: list[str] = field(default_factory=list)
    skipped_rows: int = 0


BUY_TOKENS = {"buy", "b", "买入", "買入", "bought", "purchase"}
SELL_TOKENS = {"sell", "s", "卖出", "賣出", "sold", "sale"}


def _norm_header(value: Any) -> str:
    text = str(value or "").strip().lower()
    return re.sub(r"[\s_\-./]+", "", text)


def _find_col(headers: list[str], aliases: Iterable[str]) -> int:
    normalized = [_norm_header(h) for h in headers]
    for alias in aliases:
        target = _norm_header(alias)
        if target in normalized:
            return normalized.index(target)
    for alias in aliases:
        target = _norm_header(alias)
        if len(target) < 2:
            continue
        for i, h in enumerate(normalized):
            if target in h or h in target:
                return i
    return -1


def _parse_number(value: Any) -> float:
    if value is None or value == "":
        return 0.0
    if isinstance(value, (int, float)):
        return float(value)
    raw = str(value).strip().strip("'\"`")
    if not raw:
        return 0.0
    negative = raw.startswith("(") and raw.endswith(")") or raw.startswith("-")
    raw = raw.replace("(", "").replace(")", "").replace(",", "")
    if re.fullmatch(r"\d+\.\d{3},\d+", raw):
        raw = raw.replace(".", "").replace(",", ".")
    elif re.fullmatch(r"\d+,\d+", raw):
        raw = raw.replace(",", ".")
    try:
        n = float(raw)
    except ValueError:
        return 0.0
    return -abs(n) if negative else n


def _excel_serial_to_date(serial: float) -> datetime:
    # Excel epoch 1899-12-30
    return datetime(1899, 12, 30) + timedelta(days=float(serial))


def _parse_date(value: Any) -> str:
    if value is None or value == "":
        return ""
    if isinstance(value, datetime):
        return value.strftime("%Y-%m-%d")
    if isinstance(value, (int, float)):
        try:
            return _excel_serial_to_date(float(value)).strftime("%Y-%m-%d")
        except (OverflowError, ValueError):
            return ""
    raw = str(value).strip()
    m = re.search(r"(\d{4})[/\-.](\d{1,2})[/\-.](\d{1,2})", raw)
    if m:
        return f"{m.group(1)}-{int(m.group(2)):02d}-{int(m.group(3)):02d}"
    for fmt in ("%Y-%m-%d", "%m/%d/%Y", "%Y/%m/%d", "%d/%m/%Y"):
        try:
            return datetime.strptime(raw[:10], fmt).strftime("%Y-%m-%d")
        except ValueError:
            continue
    return ""


def _parse_datetime(value: Any) -> str:
    if value is None or value == "":
        return ""
    if isinstance(value, datetime):
        return value.isoformat(timespec="seconds")
    if isinstance(value, (int, float)):
        try:
            return _excel_serial_to_date(float(value)).isoformat(timespec="seconds")
        except (OverflowError, ValueError):
            return ""
    raw = str(value).strip()
    m = re.search(
        r"(\d{4})[/\-.](\d{1,2})[/\-.](\d{1,2})(?:[ T](\d{1,2}):(\d{2})(?::(\d{2}))?)?",
        raw,
    )
    if m:
        hh = int(m.group(4) or 0)
        mm = int(m.group(5) or 0)
        ss = int(m.group(6) or 0)
        return (
            f"{m.group(1)}-{int(m.group(2)):02d}-{int(m.group(3)):02d}"
            f"T{hh:02d}:{mm:02d}:{ss:02d}"
        )
    return _parse_date(raw)


def _parse_side(value: Any, quantity: float | None = None) -> str | None:
    raw = str(value or "").strip().lower()
    if raw in BUY_TOKENS or "买" in raw or "買" in raw or "buy" in raw:
        return "buy"
    if raw in SELL_TOKENS or "卖" in raw or "賣" in raw or "sell" in raw:
        return "sell"
    if quantity is not None:
        if quantity > 0:
            return "buy"
        if quantity < 0:
            return "sell"
    return None


def _clean_symbol(value: Any) -> str:
    return re.sub(r"\s+", "", str(value or "").strip()).upper()


def _row_get(row: list[Any], idx: int) -> Any:
    if idx < 0 or idx >= len(row):
        return None
    return row[idx]


def detect_broker(headers: list[str], file_name: str) -> str:
    joined = " ".join(_norm_header(h) for h in headers)
    name = file_name.lower()
    if "recordtype" in joined and "tradedate" in joined and "settleddate" in joined:
        return "firstrade"
    if "firstrade" in name:
        return "firstrade"
    futu_hints = ("成交时间", "成交時間", "代码名称", "代碼名稱", "方向", "方向", "币种", "幣種")
    if any(_norm_header(h) in joined for h in futu_hints) or "交易流水" in joined:
        return "futu"
    if "futu" in name or "富途" in name or "moomoo" in name:
        return "futu"
    if "symbol" in joined and "action" in joined and "quantity" in joined:
        return "firstrade"
    return "unknown"


def _read_csv_rows(path: Path) -> list[list[Any]]:
    text = path.read_text(encoding="utf-8-sig", errors="replace")
    return list(csv.reader(text.splitlines()))


def _read_xlsx_tables(path: Path) -> list[tuple[str, list[list[Any]]]]:
    from openpyxl import load_workbook

    wb = load_workbook(path, data_only=True, read_only=True)
    tables: list[tuple[str, list[list[Any]]]] = []
    for sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
        rows: list[list[Any]] = []
        for row in ws.iter_rows(values_only=True):
            rows.append(list(row))
        tables.append((sheet_name, rows))
    return tables


def _parse_firstrade_rows(
    rows: list[list[Any]], file_name: str, source_label: str
) -> ParseResult:
    result = ParseResult(broker="firstrade", file_name=file_name)
    if not rows:
        result.warnings.append(f"{source_label}: 空檔")
        return result

    header_idx = 0
    for i, row in enumerate(rows[:10]):
        if any(_norm_header(c) == "symbol" for c in row):
            header_idx = i
            break
    headers = [str(c or "") for c in rows[header_idx]]
    idx = {
        "symbol": _find_col(headers, ["Symbol"]),
        "qty": _find_col(headers, ["Quantity", "Qty"]),
        "price": _find_col(headers, ["Price"]),
        "action": _find_col(headers, ["Action", "Type"]),
        "date": _find_col(headers, ["TradeDate", "Date", "Trade Date"]),
        "commission": _find_col(headers, ["Commission"]),
        "fee": _find_col(headers, ["Fee"]),
        "amount": _find_col(headers, ["Amount"]),
        "record": _find_col(headers, ["RecordType", "Record Type"]),
        "desc": _find_col(headers, ["Description"]),
    }
    if idx["symbol"] < 0 or idx["date"] < 0:
        result.warnings.append(f"{source_label}: 找不到 Firstrade 必要欄位")
        return result

    for row_no, row in enumerate(rows[header_idx + 1 :], start=header_idx + 2):
        if not any(cell not in (None, "") for cell in row):
            continue
        record = str(_row_get(row, idx["record"]) or "").strip().lower()
        if idx["record"] >= 0 and record and record != "trade":
            result.skipped_rows += 1
            continue
        symbol = _clean_symbol(_row_get(row, idx["symbol"]))
        if not symbol:
            result.skipped_rows += 1
            continue
        qty = _parse_number(_row_get(row, idx["qty"]))
        action = _parse_side(_row_get(row, idx["action"]), qty)
        if action is None:
            result.skipped_rows += 1
            continue
        qty = abs(qty)
        price = abs(_parse_number(_row_get(row, idx["price"])))
        fees = abs(_parse_number(_row_get(row, idx["commission"]))) + abs(
            _parse_number(_row_get(row, idx["fee"]))
        )
        amount = abs(_parse_number(_row_get(row, idx["amount"])))
        if amount <= 0 and qty and price:
            amount = qty * price
        date = _parse_date(_row_get(row, idx["date"]))
        if not date:
            result.warnings.append(f"{source_label} 第 {row_no} 列日期無法解析")
            result.skipped_rows += 1
            continue
        result.trades.append(
            Trade(
                broker="firstrade",
                symbol=symbol,
                side=action,
                quantity=qty,
                price=price,
                trade_date=date,
                fees=fees,
                amount=amount,
                currency="USD",
                market="US",
                source_file=file_name,
                source_row=row_no,
            )
        )
    return result


def _parse_futu_rows(
    rows: list[list[Any]], file_name: str, source_label: str
) -> ParseResult:
    result = ParseResult(broker="futu", file_name=file_name)
    if not rows:
        result.warnings.append(f"{source_label}: 空檔")
        return result

    header_idx = -1
    for i, row in enumerate(rows[:30]):
        headers = [str(c or "") for c in row]
        if _find_col(headers, ["方向", "Side", "trd_side", "买卖"]) >= 0 and _find_col(
            headers, ["代码名称", "代碼名稱", "代码", "代號", "Symbol", "code"]
        ) >= 0:
            header_idx = i
            break
    if header_idx < 0:
        result.warnings.append(f"{source_label}: 找不到富途交易欄位（需含方向／代碼）")
        return result

    headers = [str(c or "") for c in rows[header_idx]]
    idx = {
        "time": _find_col(headers, ["成交时间", "成交時間", "create_time", "时间", "時間"]),
        "date": _find_col(headers, ["交收日期", "日期", "Date"]),
        "symbol": _find_col(
            headers, ["代码名称", "代碼名稱", "代码", "代號", "Symbol", "code", "股票代码"]
        ),
        "side": _find_col(headers, ["方向", "Side", "trd_side", "买卖", "買賣"]),
        "qty": _find_col(headers, ["数量", "數量", "qty", "Quantity"]),
        "price": _find_col(headers, ["价格", "價格", "成交价", "成交價", "price", "Price"]),
        "currency": _find_col(headers, ["币种", "幣種", "currency", "Currency"]),
        "gross": _find_col(headers, ["成交金额", "成交金額", "amount"]),
        "fees": _find_col(headers, ["总费用", "總費用", "费用", "費用", "Commission"]),
        "other_fees": _find_col(headers, ["其他费用", "其他費用"]),
        "market": _find_col(headers, ["市场", "市場", "交易所", "deal_market"]),
    }

    for row_no, row in enumerate(rows[header_idx + 1 :], start=header_idx + 2):
        if not any(cell not in (None, "") for cell in row):
            continue
        symbol = _clean_symbol(_row_get(row, idx["symbol"]))
        if not symbol:
            result.skipped_rows += 1
            continue
        qty = abs(_parse_number(_row_get(row, idx["qty"])))
        side = _parse_side(_row_get(row, idx["side"]), None)
        if side is None:
            result.skipped_rows += 1
            continue
        price = abs(_parse_number(_row_get(row, idx["price"])))
        fees = abs(_parse_number(_row_get(row, idx["fees"]))) + abs(
            _parse_number(_row_get(row, idx["other_fees"]))
        )
        amount = abs(_parse_number(_row_get(row, idx["gross"])))
        if amount <= 0 and qty and price:
            amount = qty * price
        trade_time = _parse_datetime(_row_get(row, idx["time"]))
        date = trade_time[:10] if trade_time else _parse_date(_row_get(row, idx["date"]))
        if not date:
            result.warnings.append(f"{source_label} 第 {row_no} 列日期無法解析")
            result.skipped_rows += 1
            continue
        currency = str(_row_get(row, idx["currency"]) or "USD").strip().upper() or "USD"
        market = str(_row_get(row, idx["market"]) or "").strip().upper()
        result.trades.append(
            Trade(
                broker="futu",
                symbol=symbol,
                side=side,
                quantity=qty,
                price=price,
                trade_date=date,
                trade_time=trade_time,
                fees=fees,
                amount=amount,
                currency=currency,
                market=market,
                source_file=file_name,
                source_row=row_no,
            )
        )
    return result


def parse_file(path: Path, broker_hint: str | None = None) -> ParseResult:
    suffix = path.suffix.lower()
    if suffix == ".csv":
        rows = _read_csv_rows(path)
        headers = [str(c or "") for c in rows[0]] if rows else []
        broker = broker_hint or detect_broker(headers, path.name)
        if broker == "firstrade":
            return _parse_firstrade_rows(rows, path.name, path.name)
        if broker == "futu":
            return _parse_futu_rows(rows, path.name, path.name)
        # fallback: try both
        futu = _parse_futu_rows(rows, path.name, path.name)
        if futu.trades:
            return futu
        first = _parse_firstrade_rows(rows, path.name, path.name)
        if first.trades:
            return first
        result = ParseResult(broker="unknown", file_name=path.name)
        result.warnings.append(f"{path.name}: 無法辨識券商格式")
        return result

    if suffix in {".xlsx", ".xlsm"}:
        tables = _read_xlsx_tables(path)
        preferred = []
        others = []
        for sheet_name, rows in tables:
            if any(key in sheet_name for key in ("交易流水", "成交", "Trade", "历史成交", "歷史成交")):
                preferred.append((sheet_name, rows))
            else:
                others.append((sheet_name, rows))
        merged = ParseResult(broker="futu", file_name=path.name)
        for sheet_name, rows in preferred + others:
            part = _parse_futu_rows(rows, path.name, f"{path.name} / {sheet_name}")
            merged.trades.extend(part.trades)
            merged.warnings.extend(part.warnings)
            merged.skipped_rows += part.skipped_rows
            if part.trades and broker_hint != "firstrade":
                merged.broker = "futu"
        if not merged.trades:
            # rare: firstrade-like xlsx
            for sheet_name, rows in tables:
                part = _parse_firstrade_rows(rows, path.name, f"{path.name} / {sheet_name}")
                if part.trades:
                    return part
            merged.warnings.append(f"{path.name}: XLSX 中找不到可解析交易列")
        return merged

    result = ParseResult(broker="unknown", file_name=path.name)
    result.warnings.append(f"{path.name}: 不支援副檔名 {suffix}（請用 CSV 或 XLSX）")
    return result


def parse_files(paths: list[Path]) -> list[ParseResult]:
    results: list[ParseResult] = []
    for path in paths:
        hint = None
        name = path.name.lower()
        parent = path.parent.name.lower()
        if "firstrade" in name or "firstrade" in parent:
            hint = "firstrade"
        elif "futu" in name or "富途" in name or "moomoo" in name or "futu" in parent:
            hint = "futu"
        results.append(parse_file(path, broker_hint=hint))
    return results
